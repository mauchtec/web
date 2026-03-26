import csv
import io

from django.db import IntegrityError
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.shortcuts import get_object_or_404
from rest_framework.decorators import api_view, permission_classes, parser_classes
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework import status

from GateCore.models import AuditTrail, Occupancy, Person, Property, Site, Unit
from GateCore.api_views.device_sync import mark_device_sync_pending


def _require_staff(request):
    if not request.user.is_staff:
        return Response({"detail": "Forbidden"}, status=status.HTTP_403_FORBIDDEN)
    return None


def _split_full_name(full_name: str):
    parts = (full_name or "").strip().split()
    if not parts:
        return "", ""
    return parts[0], " ".join(parts[1:])


def _parse_tags(raw_tags: str):
    if not raw_tags:
        return []
    return [tag.strip() for tag in raw_tags.split(",") if tag.strip()]


def _parse_bool(value, default=False):
    if value in (True, "true", "True", "1", 1, "on"):
        return True
    if value in (False, "false", "False", "0", 0, "off"):
        return False
    return default


def _parse_date(value):
    if not value:
        return None
    if hasattr(value, "isoformat"):
        return value
    return parse_date(str(value))


def _parse_datetime(value):
    if not value:
        return None
    if hasattr(value, "isoformat"):
        return value
    value = str(value).strip()
    if not value:
        return None
    try:
        from django.utils.dateparse import parse_datetime
        return parse_datetime(value)
    except Exception:
        return None


def _parse_decimal(value):
    if value in (None, ""):
        return None
    try:
        from decimal import Decimal
        return Decimal(str(value))
    except Exception:
        return None


def _parse_int(value):
    if value in (None, ""):
        return None
    try:
        return int(value)
    except Exception:
        return None


def _parse_excel_bool(value, default=False):
    if value in (None, ""):
        return default
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text in {"1", "true", "yes", "y", "on"}:
        return True
    if text in {"0", "false", "no", "n", "off"}:
        return False
    return default


def _import_file_rows(upload):
    name = str(getattr(upload, "name", "") or "").lower()
    if name.endswith(".csv"):
        text = upload.read().decode("utf-8-sig")
        return list(csv.DictReader(io.StringIO(text)))
    if name.endswith(".xls"):
        import xlrd

        workbook = xlrd.open_workbook(file_contents=upload.read())
        sheet = workbook.sheet_by_index(0)
        headers = [str(sheet.cell_value(0, idx)).strip() for idx in range(sheet.ncols)]
        rows = []
        for row_idx in range(1, sheet.nrows):
            row = {}
            for col_idx, header in enumerate(headers):
                row[header] = sheet.cell_value(row_idx, col_idx)
            rows.append(row)
        return rows
    if name.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("XLSX import requires openpyxl. Install it or save the file as CSV/XLS.") from exc

        workbook = load_workbook(upload, read_only=True, data_only=True)
        sheet = workbook.active
        rows_iter = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows_iter, [])]
        rows = []
        for values in rows_iter:
            row = {}
            for idx, header in enumerate(headers):
                row[header] = values[idx] if idx < len(values) else None
            rows.append(row)
        return rows
    raise ValueError("Unsupported file type. Upload CSV, XLS, or XLSX.")


def _clean_import_value(row, key):
    value = row.get(key)
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


@api_view(["POST"])
@permission_classes([IsAuthenticated])
@parser_classes([MultiPartParser, FormParser])
def import_units_excel(request):
    forbidden = _require_staff(request)
    if forbidden:
        return forbidden

    upload = request.FILES.get("file")
    if not upload:
        return Response({"success": False, "error": "Upload a file first."}, status=status.HTTP_400_BAD_REQUEST)

    try:
        raw_rows = _import_file_rows(upload)
    except Exception as exc:
        return Response({"success": False, "error": str(exc)}, status=status.HTTP_400_BAD_REQUEST)

    created_properties = 0
    updated_properties = 0
    created_units = 0
    updated_units = 0
    errors = []

    for index, row in enumerate(raw_rows, start=2):
        site_name = _clean_import_value(row, "site_name")
        property_name = _clean_import_value(row, "property_name")
        unit_code = _clean_import_value(row, "unit_code")
        if not site_name or not property_name or not unit_code:
            errors.append(f"Row {index}: site_name, property_name and unit_code are required.")
            continue

        site = Site.objects.filter(name__iexact=site_name).first()
        if not site:
            errors.append(f"Row {index}: site '{site_name}' was not found.")
            continue

        property_defaults = {
            "property_type": _clean_import_value(row, "property_type") or Property.PROPERTY_TYPES[0][0],
            "property_code": _clean_import_value(row, "property_code") or None,
            "address": _clean_import_value(row, "property_address"),
            "location": _clean_import_value(row, "property_location"),
            "total_floors": _parse_int(row.get("total_floors")) or 1,
        }
        if property_defaults["property_type"] not in dict(Property.PROPERTY_TYPES):
            errors.append(f"Row {index}: invalid property_type '{property_defaults['property_type']}'.")
            continue

        property_obj, property_created = Property.objects.get_or_create(
            site=site,
            name=property_name,
            defaults=property_defaults,
        )
        if property_created:
            created_properties += 1
        else:
            changed = False
            for field, value in property_defaults.items():
                if value not in ("", None) and getattr(property_obj, field) != value:
                    setattr(property_obj, field, value)
                    changed = True
            if changed:
                property_obj.save()
                updated_properties += 1

        unit_defaults = {
            "unit_type": _clean_import_value(row, "unit_type") or Unit.UNIT_TYPES[0][0],
            "floor": _parse_int(row.get("floor")),
            "status": _clean_import_value(row, "status") or Unit.STATUS_CHOICES[0][0],
            "area_sqft": _parse_decimal(row.get("area_sqft")),
            "bedrooms": _parse_int(row.get("bedrooms")) or 0,
            "bathrooms": _parse_int(row.get("bathrooms")) or 1,
            "access_code": _clean_import_value(row, "access_code") or None,
            "has_parking": _parse_excel_bool(row.get("has_parking"), default=False),
            "parking_spots": _parse_int(row.get("parking_spots")) or 0,
            "permissions": _clean_import_value(row, "permissions") or Unit.PERMISSIONS_CHOICES[0][0],
            "is_active": _parse_excel_bool(row.get("is_active"), default=True),
        }
        if unit_defaults["unit_type"] not in dict(Unit.UNIT_TYPES):
            errors.append(f"Row {index}: invalid unit_type '{unit_defaults['unit_type']}'.")
            continue
        if unit_defaults["status"] not in dict(Unit.STATUS_CHOICES):
            errors.append(f"Row {index}: invalid status '{unit_defaults['status']}'.")
            continue
        if unit_defaults["permissions"] not in dict(Unit.PERMISSIONS_CHOICES):
            errors.append(f"Row {index}: invalid permissions '{unit_defaults['permissions']}'.")
            continue

        unit, unit_created = Unit.objects.get_or_create(
            property=property_obj,
            unit_code=unit_code,
            defaults=unit_defaults,
        )
        if unit_created:
            created_units += 1
        else:
            changed = False
            for field, value in unit_defaults.items():
                if value != getattr(unit, field):
                    setattr(unit, field, value)
                    changed = True
            if changed:
                unit.save()
                updated_units += 1

    AuditTrail.objects.create(
        action="import_units_excel",
        performed_by=request.user if request.user.is_authenticated else None,
        details=f"Imported units file: properties created={created_properties}, properties updated={updated_properties}, units created={created_units}, units updated={updated_units}, errors={len(errors)}",
    )
    if created_properties or updated_properties or created_units or updated_units:
        mark_device_sync_pending("Units imported")

    return Response(
        {
            "success": True,
            "created_properties": created_properties,
            "updated_properties": updated_properties,
            "created_units": created_units,
            "updated_units": updated_units,
            "errors": errors,
        }
    )


def _serialize_unit(unit):
    property_obj = unit.property
    site = property_obj.site if property_obj else None
    return {
        "id": str(unit.id),
        "site": site.name if site else "",
        "property": property_obj.name if property_obj else "",
        "unit_code": unit.unit_code,
        "unit_type": unit.unit_type,
        "unit_type_display": unit.get_unit_type_display(),
        "status": unit.status,
        "status_display": unit.get_status_display(),
        "permissions": unit.permissions,
        "permissions_display": unit.get_permissions_display(),
        "is_active": bool(unit.is_active),
    }


def _serialize_person(person):
    return {
        "id": str(person.id),
        "full_name": person.full_name,
        "id_number": person.id_number or "",
        "card_number": person.card_number or "",
        "phone": person.phone or "",
        "phone_device_type": person.phone_device_type or "",
        "phone_otp": person.phone_otp or "",
        "facial_recognition_enabled": bool(person.facial_recognition_enabled),
        "face_provider": person.face_provider or "",
        "face_subject_id": person.face_subject_id or "",
        "face_enrollment_status": person.face_enrollment_status or "not_started",
        "face_enrollment_status_display": person.get_face_enrollment_status_display() if person.face_enrollment_status else "",
        "face_enrollment_quality_score": str(person.face_enrollment_quality_score) if person.face_enrollment_quality_score is not None else "",
        "photo_quality_score": str(person.face_enrollment_quality_score) if person.face_enrollment_quality_score is not None else "",
        "face_enrolled_at": person.face_enrolled_at.isoformat() if person.face_enrolled_at else "",
        "face_last_synced_at": person.face_last_synced_at.isoformat() if person.face_last_synced_at else "",
        "face_enrollment_error": person.face_enrollment_error or "",
        "face_reference_image_id": person.face_reference_image_id or "",
        "face_user_type": person.face_user_type if person.face_user_type is not None else "",
        "face_access_password": person.face_access_password or "",
        "face_pass_rule_id": person.face_pass_rule_id or "",
        "face_tts_name": person.face_tts_name or "",
        "face_effective_from": person.face_effective_from.isoformat() if person.face_effective_from else "",
        "face_max_pass_count": person.face_max_pass_count if person.face_max_pass_count is not None else "",
        "face_pass_count_cycle": person.face_pass_count_cycle if person.face_pass_count_cycle is not None else "",
        "photo_review_status": person.photo_review_status or "not_started",
        "photo_review_status_display": person.get_photo_review_status_display() if person.photo_review_status else "",
        "photo_similarity_score": str(person.photo_similarity_score) if person.photo_similarity_score is not None else "",
        "photo_reviewed_at": person.photo_reviewed_at.isoformat() if person.photo_reviewed_at else "",
        "photo_review_error": person.photo_review_error or "",
        "photo_review_attempt_url": getattr(person.photo_review_attempt, "url", "") if getattr(person, "photo_review_attempt", None) else "",
        "email": person.email or "",
        "gender": person.gender or "",
        "gender_display": person.get_gender_display() if person.gender else "",
        "date_of_birth": person.date_of_birth.isoformat() if person.date_of_birth else "",
        "access_expires_on": person.access_expires_on.isoformat() if person.access_expires_on else "",
        "emergency_contact_name": person.emergency_contact_name or "",
        "emergency_contact_phone": person.emergency_contact_phone or "",
        "notes": person.notes or "",
        "tags": person.tags or [],
        "device_identifier": person.device_identifier or "",
        "device_name": person.device_name or "",
        "device_model": person.device_model or "",
        "device_serial_number": person.device_serial_number or "",
        "device_os": person.device_os or "",
        "device_app_version": person.device_app_version or "",
    }


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def add_unit(request):
    forbidden = _require_staff(request)
    if forbidden:
        return forbidden

    data = request.data
    site_name = (data.get("site") or "").strip()
    property_name = (data.get("property") or "").strip()
    unit_code = (data.get("unit_code") or "").strip()
    unit_type = (data.get("unit_type") or "").strip()
    unit_status = (data.get("status") or "").strip()
    permissions = (data.get("permissions") or "").strip()
    is_active = _parse_bool(data.get("is_active"), default=False)

    if not site_name or not property_name or not unit_code:
        return Response({"success": False, "error": "Site, property, and unit code are required."}, status=status.HTTP_400_BAD_REQUEST)

    site_qs = Site.objects.filter(name__iexact=site_name)
    if not site_qs.exists():
        return Response({"success": False, "error": "Site not found. Create the site first."}, status=status.HTTP_404_NOT_FOUND)
    if site_qs.count() > 1:
        return Response({"success": False, "error": "Multiple sites match this name. Please use a unique site name."}, status=status.HTTP_400_BAD_REQUEST)
    site = site_qs.first()
    if site.site_type != "estate":
        return Response({"success": False, "error": "Only estate sites are supported on this page."}, status=status.HTTP_400_BAD_REQUEST)

    property_qs = Property.objects.filter(site=site, name__iexact=property_name)
    if not property_qs.exists():
        return Response({"success": False, "error": "Property not found for this site. Create the property first."}, status=status.HTTP_404_NOT_FOUND)
    if property_qs.count() > 1:
        return Response({"success": False, "error": "Multiple properties match this name. Please use a unique property name."}, status=status.HTTP_400_BAD_REQUEST)
    property_obj = property_qs.first()

    if unit_type and unit_type not in dict(Unit.UNIT_TYPES):
        return Response({"success": False, "error": "Invalid unit type."}, status=status.HTTP_400_BAD_REQUEST)
    if unit_status and unit_status not in dict(Unit.STATUS_CHOICES):
        return Response({"success": False, "error": "Invalid status."}, status=status.HTTP_400_BAD_REQUEST)
    if permissions and permissions not in dict(Unit.PERMISSIONS_CHOICES):
        return Response({"success": False, "error": "Invalid permissions."}, status=status.HTTP_400_BAD_REQUEST)
    if not permissions:
        permissions = Unit.PERMISSIONS_CHOICES[0][0]

    if Unit.objects.filter(property=property_obj, unit_code__iexact=unit_code).exists():
        return Response({"success": False, "error": "Unit code already exists for this property."}, status=status.HTTP_400_BAD_REQUEST)

    try:
        unit = Unit.objects.create(
            property=property_obj,
            unit_code=unit_code,
            unit_type=unit_type or Unit.UNIT_TYPES[0][0],
            status=unit_status or Unit.STATUS_CHOICES[0][0],
            permissions=permissions,
            is_active=is_active,
        )
        AuditTrail.objects.create(
            unit=unit,
            action="create_unit",
            performed_by=request.user if request.user.is_authenticated else None,
            details=f"Created unit {unit}",
        )
        mark_device_sync_pending("Unit created")
    except IntegrityError as exc:
        return Response({"success": False, "error": str(exc)}, status=status.HTTP_400_BAD_REQUEST)
    except Exception as exc:
        return Response({"success": False, "error": str(exc)}, status=status.HTTP_400_BAD_REQUEST)

    return Response({"success": True, "unit": _serialize_unit(unit)})


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def update_unit(request, unit_id):
    forbidden = _require_staff(request)
    if forbidden:
        return forbidden

    unit = get_object_or_404(Unit, id=unit_id)
    data = request.data

    try:
        if "unit_code" in data:
            unit.unit_code = data.get("unit_code") or unit.unit_code
        if "unit_type" in data and data.get("unit_type"):
            unit.unit_type = data.get("unit_type")
        if "status" in data and data.get("status"):
            unit.status = data.get("status")
        if "permissions" in data and data.get("permissions"):
            unit.permissions = data.get("permissions")
        if "is_active" in data:
            unit.is_active = _parse_bool(data.get("is_active"), default=False)
        else:
            unit.is_active = False
        unit.save()
        AuditTrail.objects.create(
            unit=unit,
            action="update_unit",
            performed_by=request.user if request.user.is_authenticated else None,
            details=f"Edited unit {unit}",
        )
        mark_device_sync_pending("Unit updated")
        return Response({"success": True, "unit": _serialize_unit(unit)})
    except Exception as exc:
        return Response({"success": False, "error": str(exc)}, status=status.HTTP_400_BAD_REQUEST)


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def add_person(request):
    forbidden = _require_staff(request)
    if forbidden:
        return forbidden

    data = request.data
    unit_id = data.get("unit_id") or data.get("unit")
    unit = get_object_or_404(Unit, id=unit_id)

    try:
        first_name, last_name = _split_full_name(data.get("full_name", ""))
        person = Person.objects.create(
            first_name=first_name,
            last_name=last_name,
            id_number=data.get("id_number"),
            card_number=data.get("card_number", ""),
            phone=data.get("phone"),
            phone_device_type=data.get("phone_device_type"),
            phone_otp=data.get("phone_otp"),
            facial_recognition_enabled=(data.get("facial_recognition_enabled") == "Yes"),
            email=data.get("email"),
            date_of_birth=_parse_date(data.get("date_of_birth")),
            access_expires_on=_parse_date(data.get("access_expires_on")),
            face_provider=(data.get("face_provider") or "compreface").strip(),
            face_subject_id=(data.get("face_subject_id") or "").strip(),
            face_enrollment_status=(data.get("face_enrollment_status") or "not_started").strip(),
            face_enrollment_quality_score=_parse_decimal(data.get("face_enrollment_quality_score")),
            face_enrolled_at=_parse_datetime(data.get("face_enrolled_at")),
            face_last_synced_at=_parse_datetime(data.get("face_last_synced_at")),
            face_enrollment_error=data.get("face_enrollment_error", ""),
            face_user_type=_parse_int(data.get("face_user_type")),
            face_access_password=data.get("face_access_password", ""),
            face_pass_rule_id=data.get("face_pass_rule_id", ""),
            face_tts_name=data.get("face_tts_name", ""),
            face_effective_from=_parse_datetime(data.get("face_effective_from")),
            face_max_pass_count=_parse_int(data.get("face_max_pass_count")),
            face_pass_count_cycle=_parse_int(data.get("face_pass_count_cycle")),
            emergency_contact_name=data.get("emergency_contact", ""),
            tags=_parse_tags(data.get("tags")),
            device_identifier=data.get("device_identifier", ""),
            device_name=data.get("device_name", ""),
            device_model=data.get("device_model", ""),
            device_serial_number=data.get("device_serial_number", ""),
            device_os=data.get("device_os", ""),
            device_app_version=data.get("device_app_version", ""),
        )
        AuditTrail.objects.create(
            person=person,
            action="create",
            performed_by=request.user if request.user.is_authenticated else None,
            details=f"Created person {person} in unit {unit}",
        )
        mark_device_sync_pending("Resident created")
        start_date = _parse_date(data.get("start_date")) or timezone.now().date()
        Occupancy.objects.create(
            person=person,
            unit=unit,
            role=data.get("role", "tenant"),
            start_date=start_date,
        )
        return Response({"success": True, "person_id": str(person.id), "person": _serialize_person(person)})
    except IntegrityError as exc:
        if "UNIQUE constraint failed" in str(exc):
            return Response({"success": False, "error": "ID Number must be unique."}, status=status.HTTP_400_BAD_REQUEST)
        return Response({"success": False, "error": str(exc)}, status=status.HTTP_400_BAD_REQUEST)
    except Exception as exc:
        return Response({"success": False, "error": str(exc)}, status=status.HTTP_400_BAD_REQUEST)


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def update_person(request, person_id):
    forbidden = _require_staff(request)
    if forbidden:
        return forbidden

    person = get_object_or_404(Person, id=person_id)
    data = request.data

    try:
        first_name, last_name = _split_full_name(data.get("full_name", ""))
        person.first_name = first_name
        person.last_name = last_name
        person.id_number = data.get("id_number")
        person.card_number = data.get("card_number", "")
        person.phone = data.get("phone")
        person.phone_device_type = data.get("phone_device_type")
        person.phone_otp = data.get("phone_otp")
        person.facial_recognition_enabled = data.get("facial_recognition_enabled") == "Yes"
        person.email = data.get("email")
        person.gender = data.get("gender") or ""
        person.date_of_birth = _parse_date(data.get("date_of_birth"))
        person.access_expires_on = _parse_date(data.get("access_expires_on"))
        person.face_provider = (data.get("face_provider") or person.face_provider or "compreface").strip()
        person.face_subject_id = (data.get("face_subject_id") or person.face_subject_id or "").strip()
        person.face_enrollment_status = (data.get("face_enrollment_status") or person.face_enrollment_status or "not_started").strip()
        if "face_enrollment_quality_score" in data:
            person.face_enrollment_quality_score = _parse_decimal(data.get("face_enrollment_quality_score"))
        if "face_enrolled_at" in data:
            person.face_enrolled_at = _parse_datetime(data.get("face_enrolled_at"))
        if "face_last_synced_at" in data:
            person.face_last_synced_at = _parse_datetime(data.get("face_last_synced_at"))
        if "face_enrollment_error" in data:
            person.face_enrollment_error = data.get("face_enrollment_error") or ""
        if "face_user_type" in data:
            person.face_user_type = _parse_int(data.get("face_user_type"))
        if "face_access_password" in data:
            person.face_access_password = data.get("face_access_password") or ""
        if "face_pass_rule_id" in data:
            person.face_pass_rule_id = data.get("face_pass_rule_id") or ""
        if "face_tts_name" in data:
            person.face_tts_name = data.get("face_tts_name") or ""
        if "face_effective_from" in data:
            person.face_effective_from = _parse_datetime(data.get("face_effective_from"))
        if "face_max_pass_count" in data:
            person.face_max_pass_count = _parse_int(data.get("face_max_pass_count"))
        if "face_pass_count_cycle" in data:
            person.face_pass_count_cycle = _parse_int(data.get("face_pass_count_cycle"))
        person.emergency_contact_name = data.get("emergency_contact_name") or data.get("emergency_contact", "")
        person.emergency_contact_phone = data.get("emergency_contact_phone", "")
        person.tags = _parse_tags(data.get("tags"))
        person.device_identifier = data.get("device_identifier", "")
        person.device_name = data.get("device_name", "")
        person.device_model = data.get("device_model", "")
        person.device_serial_number = data.get("device_serial_number", "")
        person.device_os = data.get("device_os", "")
        person.device_app_version = data.get("device_app_version", "")
        person.notes = data.get("notes", "")
        person.save()
        AuditTrail.objects.create(
            person=person,
            action="edit",
            performed_by=request.user if request.user.is_authenticated else None,
            details=f"Edited person {person}",
        )
        mark_device_sync_pending("Resident updated")
        return Response({"success": True, "person": _serialize_person(person)})
    except Exception as exc:
        return Response({"success": False, "error": str(exc)}, status=status.HTTP_400_BAD_REQUEST)


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def link_person_to_unit(request):
    forbidden = _require_staff(request)
    if forbidden:
        return forbidden

    data = request.data
    person_id = data.get("person_id")
    unit_id = data.get("unit_id")
    if not person_id or not unit_id:
        return Response({"success": False, "error": "Missing person or unit."}, status=status.HTTP_400_BAD_REQUEST)

    try:
        person = Person.objects.get(id=person_id)
        unit = Unit.objects.get(id=unit_id)
        if Occupancy.objects.filter(person=person, unit=unit, is_active=True, is_deleted=False, end_date__isnull=True).exists():
            return Response({"success": False, "error": "Person is already linked to this unit."}, status=status.HTTP_400_BAD_REQUEST)
        Occupancy.objects.create(
            person=person,
            unit=unit,
            role="tenant",
            start_date=timezone.now().date(),
            is_active=True,
            is_deleted=False,
        )
        AuditTrail.objects.create(
            person=person,
            unit=unit,
            action="link",
            performed_by=request.user if request.user.is_authenticated else None,
            details=f"Linked {person} to unit {unit}",
        )
        mark_device_sync_pending("Resident linked to unit")
        return Response({"success": True})
    except Person.DoesNotExist:
        return Response({"success": False, "error": "Person not found."}, status=status.HTTP_404_NOT_FOUND)
    except Unit.DoesNotExist:
        return Response({"success": False, "error": "Unit not found."}, status=status.HTTP_404_NOT_FOUND)
    except Exception as exc:
        return Response({"success": False, "error": str(exc)}, status=status.HTTP_400_BAD_REQUEST)


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def reallocate_person(request):
    forbidden = _require_staff(request)
    if forbidden:
        return forbidden

    data = request.data
    person_id = data.get("person_id")
    new_unit_id = data.get("new_unit_id") or data.get("unit_id")
    if not person_id or not new_unit_id:
        return Response({"success": False, "error": "Missing person or unit."}, status=status.HTTP_400_BAD_REQUEST)

    try:
        person = Person.objects.get(id=person_id)
        new_unit = Unit.objects.get(id=new_unit_id)

        Occupancy.objects.filter(
            person=person,
            is_active=True,
            is_deleted=False,
            end_date__isnull=True,
        ).update(is_active=False, end_date=timezone.now().date())

        existing_occupancy = Occupancy.objects.filter(
            person=person,
            unit=new_unit,
            is_deleted=False,
        ).first()
        AuditTrail.objects.create(
            person=person,
            unit=new_unit,
            action="move",
            performed_by=request.user if request.user.is_authenticated else None,
            details=f"Moved {person} to unit {new_unit}",
        )
        mark_device_sync_pending("Resident reallocated")
        if existing_occupancy:
            existing_occupancy.end_date = None
            existing_occupancy.is_active = True
            existing_occupancy.save()
            return Response({"success": True})

        Occupancy.objects.create(
            person=person,
            unit=new_unit,
            role="tenant",
            start_date=timezone.now().date(),
            is_active=True,
            is_deleted=False,
        )
        return Response({"success": True})
    except Person.DoesNotExist:
        return Response({"success": False, "error": "Person not found."}, status=status.HTTP_404_NOT_FOUND)
    except Unit.DoesNotExist:
        return Response({"success": False, "error": "Unit not found."}, status=status.HTTP_404_NOT_FOUND)
    except Exception as exc:
        return Response({"success": False, "error": str(exc)}, status=status.HTTP_400_BAD_REQUEST)
